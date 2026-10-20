"""The ideal data set stays ideal: no invalid internal values, no provider defects, one current record per source."""
import datetime as dt
from pathlib import Path

import openpyxl

from app.engine.values import validate_element

IDEAL = Path(__file__).resolve().parents[2] / "seed" / "ideal-files"


def test_same_files_as_the_fragmented_set():
    samples = IDEAL.parent / "sample-files"
    names = lambda d: sorted(p.relative_to(d).as_posix() for p in d.rglob("*") if p.is_file() and p.suffix in (".xlsx", ".csv"))
    assert names(IDEAL) == names(samples)


def test_internal_values_are_valid():
    for p in (IDEAL / "internal").glob("*.xlsx"):
        ws = openpyxl.load_workbook(p, read_only=True)["Data Dictionary"]
        bad = [(r[4], r[6]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and validate_element(r[5], r[6])]
        assert not bad, (p.name, bad[:5])


def test_external_records_are_clean_and_current():
    for p in (IDEAL / "external").glob("*.xlsx"):
        wb = openpyxl.load_workbook(p, read_only=True)
        rows = [r for r in wb["External Data"].iter_rows(min_row=2, values_only=True) if r[0]]
        assert rows and all(r[31] == "VALIDATED" and r[32] == "HIGH" for r in rows), p.name
        assert len({r[5] for r in rows}) == len(rows), p.name                     # one record per source
        assert all((dt.date(2026, 10, 9) - dt.date.fromisoformat(r[8])).days <= 7 for r in rows)
        assert not [r for r in wb["Data Quality Defects"].iter_rows(min_row=2, values_only=True) if any(r)]
