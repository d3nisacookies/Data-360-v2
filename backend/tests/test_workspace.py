"""Credit memo, Phase 2, Phase 3, fix plan, project plans, reports and policies through the API.
Runs after test_api.py has uploaded the sample files."""
import io
import zipfile

import openpyxl
import pytest

from .test_api import borrower_id


def act(client, h, bid, area, action, **kw):
    r = client.post(f"/api/borrowers/{bid}/{area}/action", json={"action": action, **kw}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def is_docx(content: bytes):
    return zipfile.ZipFile(io.BytesIO(content)).namelist().count("word/document.xml") == 1


@pytest.fixture(scope="module")
def bid(client, h):
    b = borrower_id(client, h, "Meridian Renewables Limited")
    # phases stay locked until Phase 1 is complete
    assert client.get(f"/api/borrowers/{b}/memo", headers=h).json()["available"] is False
    assert client.get(f"/api/borrowers/{b}/p2", headers=h).json()["available"] is False
    for s in ("collect", "seek", "clean", "assess", "normalise"):
        assert client.post(f"/api/borrowers/{b}/{s}", json={}, headers=h).status_code == 200
    return b


def test_pick_list_extras(client, h, bid):
    a = client.post(f"/api/borrowers/{bid}/assess", json={"extra": ["Materiality", "Conflict resolution", "Independence"]}, headers=h).json()["result"]
    assert set(a["crit"]) == {"Materiality", "Conflict resolution", "Independence"}
    assert all("basis" in v and v["status"] in ("high", "med", "low", "review", "none") for v in a["crit"].values())
    meta = client.get("/api/meta", headers=h).json()
    assert len([i for g in meta["picks"]["assess"] for i in g["items"]]) == 29
    # re-running Assess makes Normalise out of date, which keeps the memo locked until it runs again
    assert client.get(f"/api/borrowers/{bid}/memo", headers=h).json()["available"] is False
    n = client.post(f"/api/borrowers/{bid}/normalise", json={"extra": ["IFRS 9 definitions"]}, headers=h).json()["result"]
    assert "IFRS 9 definitions" in n["picks"]


def test_memo_build_approve_retain(client, h, bid):
    v = act(client, h, bid, "memo", "run", arg={"from": 1, "to": 5})
    assert [s["title"] for s in v["sections"]][:2] == ["5.1 Prepare", "5.2 Analyse"]
    assert v["decision"] in ("Approve", "Approve With Conditions", "Refer", "Decline")
    # approvals go in order
    r = client.post(f"/api/borrowers/{bid}/memo/action", json={"action": "approval", "arg": {"i": 2, "field": "decision", "value": "Refer"}}, headers=h)
    assert r.status_code == 409
    assert client.post(f"/api/borrowers/{bid}/memo/action", json={"action": "retain"}, headers=h).status_code == 409
    act(client, h, bid, "memo", "edit", arg={"key": "2|Legal Entity Name", "to": "Meridian Renewables Ltd", "reason": "Registry spelling"})
    for i in range(5):
        act(client, h, bid, "memo", "approval", arg={"i": i, "field": "name", "value": f"Approver {i}"})
        v = act(client, h, bid, "memo", "approval", arg={"i": i, "field": "decision", "value": "Approve with conditions"})
    docx = client.get(f"/api/borrowers/{bid}/memo.docx", headers=h)
    assert docx.status_code == 200 and is_docx(docx.content)
    assert "Meridian Renewables Ltd" in client.get(f"/api/borrowers/{bid}/memo.html", headers=h).text
    v = act(client, h, bid, "memo", "retain")
    assert v["retained"]["version"] == "v1.0" and len(v["retained"]["hash"]) == 64
    assert client.post(f"/api/borrowers/{bid}/memo/action", json={"action": "sec.add", "arg": "A"}, headers=h).status_code == 409   # locked
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/borrowers/{bid}/evidence.xlsx", headers=h).content))
    assert "A Source traceability" in wb.sheetnames and wb["Reviewer edits"].max_row == 2


def test_phase2_roadmap(client, h, bid):
    for k in ("gap", "rec", "conf", "road"):
        v = act(client, h, bid, "p2", "run", step=k)
        assert v["sections"], k
    road = v["sections"][0]
    assert road["title"] == "Roadmap"
    v2 = act(client, h, bid, "p2", "p2.assume", step="road", arg="rate", value=200)
    assert v2["sections"][0]["blocks"] != road["blocks"]
    assert is_docx(client.get(f"/api/borrowers/{bid}/report/p2.docx", headers=h).content)
    assert client.get(f"/api/borrowers/{bid}/p3", headers=h).json()["available"] is False
    v = act(client, h, bid, "p2", "p2.approve", step="road", value="Head of Data")
    assert v["approved"]["by"] == "Head of Data"


def test_phase3_stages_and_plans(client, h, bid):
    v = act(client, h, bid, "p3", "run", step="stab")
    titles = [s["title"] for s in v["sections"]]
    assert "Critical issue resolution" in titles and "Executive confidence baseline" in titles
    if v["foot"]["block"]:
        assert client.post(f"/api/borrowers/{bid}/p3/action", json={"action": "complete", "step": "stab"}, headers=h).status_code == 409
    # Standardise is locked until Stabilise is complete
    assert client.post(f"/api/borrowers/{bid}/p3/action", json={"action": "run", "step": "std"}, headers=h).status_code == 409
    act(client, h, bid, "p3", "p3.resolveAll", step="stab")
    v = act(client, h, bid, "p3", "plan", step="stab")
    assert v["plan"]["title"].startswith("Project plan · 10 Stabilise")
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/borrowers/{bid}/plan/stab.xlsx", headers=h).content))
    assert wb.sheetnames == ["Overview", "Plan", "Gantt", "Milestones", "Workstreams", "Risks"]
    act(client, h, bid, "p3", "complete", step="stab")
    v = act(client, h, bid, "p3", "run", step="std")
    cards = next(s for s in v["sections"] if s["title"] == "Trusted datasets")["blocks"][0]["items"]
    for c in cards:
        if c["btn"] and not c["btn"]["disabled"]:
            v = act(client, h, bid, "p3", "p3.certify", step="std", arg=c["btn"]["arg"])
    act(client, h, bid, "p3", "complete", step="std")
    for k in ("opt", "trans"):
        act(client, h, bid, "p3", "run", step=k)
        act(client, h, bid, "p3", "complete", step=k)
    assert is_docx(client.get(f"/api/borrowers/{bid}/report/p3.docx", headers=h).content)
    o = client.get(f"/api/borrowers/{bid}/overview", headers=h).json()
    assert o["gates"] == {"p1": True, "p2": True, "p3": True, "open": {"1": True, "2": True, "3": True}}


def test_fix_plan_maps_to_uploaded_files(client, h, bid):
    f = client.get(f"/api/borrowers/{bid}/fix", headers=h).json()
    names = [t["name"] for g in f["groups"] for t in g["tiles"]]
    assert names and all(n for n in names)
    uploads = {u["filename"].rsplit(".", 1)[0] for u in client.get("/api/uploads", headers=h).json()}
    assert set(names) <= uploads
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/borrowers/{bid}/fix.xlsx", headers=h).content))
    assert wb.sheetnames[0] == "Summary" and len(wb.sheetnames) == len(names) + 1


def test_report_and_policies(client, h, bid):
    r = client.get(f"/api/borrowers/{bid}/report", headers=h).json()
    assert r["actions"] and r["tasks"]["n"] >= 10
    assert is_docx(client.get(f"/api/borrowers/{bid}/report.docx", headers=h).content)
    p = client.get("/api/policies", headers=h).json()
    assert len(p["all"]) == 65
    assert client.put("/api/policies", json={"on": ["K1", "K5", "nonsense"]}, headers=h).json()["on"] == ["K1", "K5"]
    client.put("/api/policies", json={"on": p["all"]}, headers=h)
    assert client.post("/api/owners", json={"source": "D&B", "owner": "Data Governance Office"}, headers=h).json()["owners"]["D&B"] == "Data Governance Office"
