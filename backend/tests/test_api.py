"""End-to-end through the API with the sample files."""
import io

import openpyxl

from .conftest import SAMPLES, upload


def borrower_id(client, h, name):
    return next(b["id"] for b in client.get("/api/borrowers", headers=h).json() if b["name"] == name)


def test_health(client):
    assert client.get("/api/health").json()["ok"]


def test_auth_required(client):
    assert client.get("/api/borrowers").status_code == 401
    assert client.post("/api/auth/login", data={"username": "admin@example.com", "password": "wrong"}).status_code == 401


def test_upload_all_samples(client, h):
    files = sorted((SAMPLES / "internal").glob("*.xlsx")) + sorted((SAMPLES / "external").glob("*.xlsx"))
    res = upload(client, h, *files)
    assert all(r["status"] == "ingested" for r in res), [(r["filename"], r["status"], r.get("error")) for r in res if r["status"] != "ingested"]
    fmts = {r["filename"]: r["format"] for r in res}
    assert fmts["Risk platforms - Data Dictionary.xlsx"] == "dictionary"
    assert fmts["External - HarbourVest.xlsx"] == "records"
    names = sorted(b["name"] for b in client.get("/api/borrowers", headers=h).json())
    assert len(names) == 7 and "HarbourVest" in names


def test_duplicate_upload_is_not_reingested(client, h):
    res = upload(client, h, SAMPLES / "external" / "External - HarbourVest.xlsx")
    assert res[0].get("duplicate") is True


def test_other_layout_csv_maps_columns_and_resolves_names(client, h):
    before = len(client.get("/api/borrowers", headers=h).json())
    res = upload(client, h, SAMPLES / "other-layouts" / "Bank B vendor export - Refinitiv.csv")
    r = res[0]
    m = {c["header"]: c["field"] for c in r["mapping"]["columns"]}
    assert m["Obligor"] == "Borrower_Name" and m["Provider"] == "External_Source"
    assert m["Facility Limit"] == "Facility_Amount" and m["Drawn"] == "Outstanding_Principal"
    assert m["Credit Rating"] == "External_Rating" and m["Report Date"] == "As_of_Date"
    assert r["status"] == "ingested", r
    # "HarbourVest Inc." and "Orion ... Pte Ltd" resolve to existing borrowers
    assert len(client.get("/api/borrowers", headers=h).json()) == before


def test_template_without_borrower_needs_mapping(client, h, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Dictionary"
    ws.append(["ID", "Data source", "Category", "Sub-category", "Data element", "Data type", "Example value", "Structure", "Sensitivity"])
    ws.append(["DE-9001", "Risk platforms", "Credit risk", "Internal ratings", "Obligor rating", "Code", "5+", "Structured", "Internal"])
    p = tmp_path / "Template.xlsx"
    wb.save(p)
    r = upload(client, h, p)[0]
    assert r["status"] == "needs_mapping" and "borrower" in r["mapping"]["missing"]
    m = r["mapping"]
    m["defaults"]["borrower"] = "HarbourVest"
    ok = client.post(f"/api/uploads/{r['id']}/ingest", json={"mapping": m}, headers=h)
    assert ok.status_code == 200 and ok.json()["status"] == "ingested"
    assert client.delete(f"/api/uploads/{r['id']}", headers=h).json()["ok"]


def test_assess_clean_reassess_normalise(client, h):
    bid = borrower_id(client, h, "HarbourVest")
    a0 = client.post(f"/api/borrowers/{bid}/assess", json={}, headers=h).json()
    before = a0["result"]["overall"]["score"]
    assert a0["result"]["dims"]["Accuracy"]["issues"], "sourced data should show wrong values"
    c = client.post(f"/api/borrowers/{bid}/clean", json={}, headers=h).json()["result"]
    assert c["log"] and c["escalations"]
    assert any(e["owner"] == "Financial Crime Compliance" for e in c["escalations"])  # internal potential sanctions match
    # Assess is now marked out of date, and re-assessing uses the cleaned data
    st = client.get(f"/api/borrowers/{bid}", headers=h).json()["steps"]
    assert st["assess"]["stale"] is True
    a1 = client.post(f"/api/borrowers/{bid}/assess", json={}, headers=h).json()["result"]
    assert a1["version"].startswith("v1.1")
    assert a1["overall"]["score"] > before
    assert a1["dims"]["Accuracy"]["score"] == 100 and a1["dims"]["Accuracy"]["before"] < 100
    assert [s["key"] for s in a1["stages"]] == ["sourced", "cleaned"]
    n = client.post(f"/api/borrowers/{bid}/normalise", json={}, headers=h).json()["result"]
    metrics = {r[0]: r[4] for r in n["out"]["Metrics"]["rows"]}
    assert metrics["Facility utilisation"] == "82.5%"
    assert n["out"]["KPIs"]["rows"][0][5].endswith("/100")


def test_selected_criteria(client, h):
    bid = borrower_id(client, h, "Orion Infrastructure Holdings Pte. Ltd.")
    r = client.post(f"/api/borrowers/{bid}/assess", json={"selected": ["Accuracy", "Validity"]}, headers=h).json()["result"]
    assert r["selected"] == ["Accuracy", "Validity"]


def test_datapack_export(client, h):
    bid = borrower_id(client, h, "HarbourVest")
    r = client.get(f"/api/borrowers/{bid}/export/datapack.xlsx", headers=h)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert {"README", "Golden record", "Records", "Audit trail", "Escalations", "Assess"} <= set(wb.sheetnames)


def test_viewer_cannot_run(client, h):
    client.post("/api/users", json={"email": "viewer@example.com", "password": "viewer12345", "role": "viewer"}, headers=h)
    t = client.post("/api/auth/login", data={"username": "viewer@example.com", "password": "viewer12345"}).json()["access_token"]
    vh = {"Authorization": f"Bearer {t}"}
    bid = borrower_id(client, vh, "HarbourVest")
    assert client.post(f"/api/borrowers/{bid}/clean", json={}, headers=vh).status_code == 403
    assert client.get(f"/api/borrowers/{bid}/assess", headers=vh).status_code == 200


def test_audit_trail(client, h):
    actions = {e["action"] for e in client.get("/api/audit", headers=h).json()}
    assert {"login", "upload", "ingest", "run.assess", "run.clean", "export.datapack"} <= actions


def test_collector_and_seeker_feed_later_steps(client, h):
    bid = borrower_id(client, h, "Meridian Renewables Limited")
    meta = client.get("/api/meta", headers=h).json()
    assert "Facility summary" in meta["outputs"] and "Registration Number" in meta["attributes"]
    # Collector: everything except Email archives, four outputs
    det = client.get(f"/api/borrowers/{bid}", headers=h).json()
    srcs = [s["source"] for s in det["sources"] if s["source"] != "Email archives"]
    co = client.post(f"/api/borrowers/{bid}/collect", json={"sources": srcs, "outputs": ["Borrower profile", "Relationship history", "Facility summary", "KYC and sanctions"]}, headers=h).json()["result"]
    rel = next(c for c in co["coverage"] if c["output"] == "Relationship history")
    assert rel["status"] == "part" and "Email archives" in rel["addable"]
    assert next(c for c in co["coverage"] if c["output"] == "Facility summary")["status"] == "ok"
    # Seeker: standard attributes plus a typed one; request a registry for a missing attribute
    sk = client.post(f"/api/borrowers/{bid}/seek", json={"attributes": ["Registration Number", "Country of Incorporation", "LEI (LEI)", "Tax ID", "Beneficial Owner", "Registered owner"],
                                                         "requested": ["IRS FATCA FFI list"]}, headers=h).json()["result"]
    st = {r["attribute"]: r for r in sk["results"]}
    assert st["Registration Number"]["status"] in ("ok", "conf") and st["Registration Number"]["hits"]
    assert st["Tax ID"]["status"] == "pending"
    assert st["Registered owner"]["hits"], "a typed attribute should be found by name similarity"
    # Assess picks them up and is scoped to the Collector's sources
    a = client.post(f"/api/borrowers/{bid}/assess", json={}, headers=h).json()["result"]
    assert a["inputs"] == {"collector": True, "seeker": True, "scope": srcs}
    assert "Seeker attributes found" in a["dims"]["Completeness"]["basis"]
    assert "Collector outputs fully covered" in a["dims"]["Business relevance"]["basis"]
    assert any(i["issue"] == "Output not fully covered" for i in a["dims"]["Business relevance"]["issues"])
    # re-running the Collector makes Assess out of date
    client.post(f"/api/borrowers/{bid}/collect", json={"sources": srcs, "outputs": ["Borrower profile"]}, headers=h)
    assert client.get(f"/api/borrowers/{bid}", headers=h).json()["steps"]["assess"]["stale"] is True
