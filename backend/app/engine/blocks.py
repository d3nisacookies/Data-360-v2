"""One document model, two renderers (Word and HTML), as in the demo console.

A document is a list of blocks: {"t": "title"|"h1"|"h2"|"h3"|"p"|"callout"|"bullets"|"numbered"|
"battery"|"kpis"|"table"|"mono"|"pagebreak", ...}.
"""
import html as _html
import io

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for callers)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = "17644A"
SHADE = {"Green": "DDF0E4", "Pass": "DDF0E4", "Compliant": "DDF0E4", "Amber": "FBEBC8", "Exception": "FBEBC8", "Tight": "FBEBC8",
         "Conflict": "FBEBC8", "Red": "F6D9D6", "Fail": "F6D9D6", "Breach": "F6D9D6"}
SHC = {"Green": "sg", "Pass": "sg", "Compliant": "sg", "Amber": "sa", "Exception": "sa", "Tight": "sa", "Conflict": "sa",
       "Red": "sr", "Fail": "sr", "Breach": "sr"}

REPORT_CSS = """.rdoc{font-family:"IBM Plex Sans","Segoe UI",Arial,sans-serif;color:#1b2420;font-size:13px;line-height:1.55;max-width:860px;margin:0 auto;background:#fff;padding:48px 56px}
.rdoc .eyebrow{font-size:11px;letter-spacing:.12em;text-transform:uppercase;color:#17644a;font-weight:600}
.rdoc h1{font-size:28px;margin:6px 0 2px;letter-spacing:-.01em}
.rdoc .rsub{font-size:16px;color:#4e5b55;margin-bottom:18px}
.rdoc .rt{border-bottom:3px solid #17644a;padding-bottom:18px;margin-bottom:8px}
.rdoc table.meta{border-collapse:collapse;font-size:12.5px}
.rdoc table.meta th{text-align:left;color:#5b6862;font-weight:500;padding:2px 18px 2px 0;white-space:nowrap;vertical-align:top}
.rdoc table.meta td{padding:2px 0}
.rdoc .r1{font-size:19px;margin:30px 0 8px;color:#17644a;border-bottom:1px solid #d6ddd9;padding-bottom:4px}
.rdoc .r2{font-size:15px;margin:20px 0 6px}
.rdoc .r3{font-size:13.5px;margin:14px 0 4px}
.rdoc p{margin:6px 0}.rdoc .rm{color:#5b6862;font-size:12px}
.rdoc ul,.rdoc ol{margin:6px 0;padding-left:22px}.rdoc li{margin:3px 0}
.rdoc .rcall{border-left:4px solid #9a6200;background:#fbf3e2;padding:10px 14px;margin:10px 0;font-weight:500}
.rdoc .rcall.ok{border-color:#1d7a4c;background:#e6f4ec}
.rdoc .rk{display:grid;grid-template-columns:repeat(auto-fit,minmax(120px,1fr));gap:8px;margin:10px 0}
.rdoc .rk div{border:1px solid #d6ddd9;border-radius:6px;padding:8px 10px;display:grid}
.rdoc .rk b{font-size:20px;font-variant-numeric:tabular-nums}.rdoc .rk span{font-size:11.5px;color:#5b6862}
.rdoc .rtw{overflow-x:auto;margin:8px 0}
.rdoc .rtab{width:100%;border-collapse:collapse;font-size:12px;table-layout:fixed}
.rdoc .rtab th{background:#17644a;color:#fff;text-align:left;font-weight:600;padding:6px 8px}
.rdoc .rtab td{border-bottom:1px solid #e1e6e3;padding:5px 8px;vertical-align:top;word-break:break-word}
.rdoc .rtab tbody tr:nth-child(even) td{background:#f5f8f6}
.rdoc .rtab .r{text-align:right;font-variant-numeric:tabular-nums}
.rdoc .rtab.sm{font-size:10.5px}.rdoc .rtab.sm td,.rdoc .rtab.sm th{padding:3px 6px}
.rdoc hr.pb{border:0;border-top:1px dashed #d6ddd9;margin:36px 0}
.rdoc .rbat{margin:12px 0;display:grid;gap:6px}.rdoc .rbl{font-weight:600}
.rdoc .rbw{display:flex;align-items:center}.rdoc .rbc{flex:1;display:flex;height:30px;border:3px solid #1b2420;border-radius:6px;padding:3px;gap:3px}
.rdoc .rbc i{display:flex;align-items:center;justify-content:center;font-style:normal;color:#fff;font-weight:700;font-size:12px;border-radius:2px}
.rdoc pre.rmono{font-family:"IBM Plex Mono",Consolas,monospace;font-size:11.5px;background:#f5f8f6;border:1px solid #d6ddd9;border-radius:6px;padding:10px 12px;overflow-x:auto;white-space:pre;margin:8px 0}
.rdoc .rtab td.sg{background:#ddf0e4!important}.rdoc .rtab td.sa{background:#fbebc8!important}.rdoc .rtab td.sr{background:#f6d9d6!important}
.rdoc .rbc i.g{background:#1d7a4c}.rdoc .rbc i.r{background:#b3362c}.rdoc .rbn{width:7px;height:13px;background:#1b2420;border-radius:0 3px 3px 0}"""


def _e(v) -> str:
    return _html.escape("" if v is None else str(v))


def to_html(B: list[dict]) -> str:
    out = []
    for b in B:
        t = b["t"]
        if t == "mono":
            out.append(f'<pre class="rmono">{_e(chr(10).join(b["lines"]))}</pre>')
        elif t == "title":
            meta = "".join(f"<tr><th>{_e(k)}</th><td>{_e(v)}</td></tr>" for k, v in b["meta"])
            out.append(f'<header class="rt"><div class="eyebrow">{_e(b.get("kicker", "Data360 · Report"))}</div><h1>{_e(b["title"])}</h1>'
                       f'<div class="rsub">{_e(b.get("sub", ""))}</div><table class="meta">{meta}</table></header>')
        elif t == "h1":
            out.append(f'<h2 class="r1">{_e(b["text"])}</h2>')
        elif t == "h2":
            out.append(f'<h3 class="r2">{_e(b["text"])}</h3>')
        elif t == "h3":
            out.append(f'<h4 class="r3">{_e(b["text"])}</h4>')
        elif t == "p":
            out.append(f'<p class="{"rm" if b.get("muted") else ""}">{_e(b["text"])}</p>')
        elif t == "callout":
            out.append(f'<div class="rcall {b.get("tone", "")}">{_e(b["text"])}</div>')
        elif t in ("bullets", "numbered"):
            tag = "ul" if t == "bullets" else "ol"
            out.append(f"<{tag}>" + "".join(f"<li>{_e(x)}</li>" for x in b["items"]) + f"</{tag}>")
        elif t == "battery":
            p = b["pct"]
            st = " → ".join(f'{_e(x["label"])} {x["pct"]}%' for x in b.get("stages") or [])
            bd = ""
            if b.get("breakdown"):
                h = b["breakdown"]
                bd = ('<div class="rtw"><table class="rtab sm"><thead><tr>' + "".join(f"<th>{_e(x)}</th>" for x in h["head"]) + "</tr></thead><tbody>"
                      + "".join("<tr>" + "".join(f"<td>{_e(v)}</td>" for v in r) + "</tr>" for r in h["rows"]) + "</tbody></table></div>")
            out.append(f'<div class="rbat"><div class="rbl">{_e(b["label"])}</div><div class="rbw"><div class="rbc"><i class="g" style="width:{p}%">{f"{p}%" if p >= 12 else ""}</i>'
                       f'<i class="r" style="width:{100 - p}%">{f"{100 - p}%" if 100 - p >= 12 else ""}</i></div><span class="rbn"></span></div>'
                       + (f'<div class="rm">{st}</div>' if st and len(b.get("stages") or []) > 1 else "") + f'<div class="rm">{_e(b.get("note", ""))}</div>{bd}</div>')
        elif t == "kpis":
            out.append('<div class="rk">' + "".join(f"<div><b>{_e(v)}</b><span>{_e(k)}</span></div>" for v, k in b["items"]) + "</div>")
        elif t == "table":
            num = set(b.get("num") or [])
            widths = b.get("widths") or [100 / max(1, len(b["head"]))] * len(b["head"])
            cols = "".join(f'<col style="width:{w}%">' for w in widths)
            head = "".join(f'<th{" class=r" if i in num else ""}>{_e(h)}</th>' for i, h in enumerate(b["head"]))
            rows = "".join("<tr>" + "".join(f'<td class="{"r " if i in num else ""}{SHC.get(str(v), "")}">{_e(v)}</td>' for i, v in enumerate(r)) + "</tr>"
                           for r in b["rows"])
            out.append(f'<div class="rtw"><table class="rtab{" sm" if b.get("small") else ""}"><colgroup>{cols}</colgroup><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></div>')
        elif t == "pagebreak":
            out.append('<hr class="pb">')
    return "\n".join(out)


def html_page(B: list[dict], title: str) -> str:
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f"<title>{_e(title)}</title><style>body{{margin:0;background:#eef2f0}}{REPORT_CSS}@media print{{body{{background:#fff}}.rdoc{{padding:0}}}}</style>"
            f'</head><body><article class="rdoc">{to_html(B)}</article></body></html>')


# ---------------- Word ----------------
def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D6DDD9")
        b.append(el)
    tblPr.append(b)


def _run(par, text, size=10, bold=False, color=None, italic=False, font="Calibri"):
    r = par.add_run("" if text is None else str(text))
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _para(doc, text, size=10, bold=False, color=None, italic=False, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    _run(p, text, size, bold, color, italic)
    return p


def _table(doc, head, rows, widths=None, small=False, num=None):
    num = set(num or [])
    n = len(head)
    widths = widths or [100 / n] * n
    total = 17.0  # cm of text width on A4 with 2 cm margins
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    sz = 7.5 if small else 8.5
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.width = Cm(total * widths[i] / 100)
        _shade(c, GREEN)
        p = c.paragraphs[0]
        _run(p, h, sz, True, "FFFFFF")
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            c = cells[i]
            c.width = Cm(total * widths[i] / 100)
            sv = "" if v is None else str(v)
            if sv in SHADE:
                _shade(c, SHADE[sv])
            elif ri % 2 == 1:
                _shade(c, "F5F8F6")
            p = c.paragraphs[0]
            if i in num:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(p, sv, sz, sv in SHADE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def to_docx(B: list[dict], title: str, header: str = "", footer: str = "Confidential · Page ") -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2))
    doc.core_properties.title = title
    doc.core_properties.author = "Data360"
    if header:
        _run(sec.header.paragraphs[0], header, 8, color="5B6862")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(fp, footer, 8, color="5B6862")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    for b in B:
        t = b["t"]
        if t == "title":
            _para(doc, (b.get("kicker") or "Data360 · Report").upper(), 9, True, GREEN, after=3)
            _para(doc, b["title"], 24, True, after=2)
            _para(doc, b.get("sub", ""), 14, color="4E5B55", after=12)
            _table(doc, b.get("metaHead") or ["Item", "Detail"], b["meta"], [34, 66])
        elif t == "h1":
            p = doc.add_heading(level=1)
            _run(p, b["text"], 15, True, GREEN)
        elif t == "h2":
            p = doc.add_heading(level=2)
            _run(p, b["text"], 12, True, "1B2420")
        elif t == "h3":
            p = doc.add_heading(level=3)
            _run(p, b["text"], 10.5, True, "1B2420")
        elif t == "p":
            _para(doc, b["text"], 9 if b.get("muted") else 10, color="5B6862" if b.get("muted") else None, italic=bool(b.get("muted")))
        elif t == "callout":
            tb = doc.add_table(rows=1, cols=1)
            c = tb.rows[0].cells[0]
            _shade(c, "E6F4EC" if b.get("tone") == "ok" else "FBF3E2")
            _run(c.paragraphs[0], b["text"], 10, True)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "bullets":
            for x in b["items"]:
                p = doc.add_paragraph(style="List Bullet")
                _run(p, x)
        elif t == "numbered":
            for i, x in enumerate(b["items"], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.5)
                p.paragraph_format.space_after = Pt(3)
                _run(p, f"{i}.  {x}")
        elif t == "battery":
            rows = [[x["label"], f'{x["pct"]}%'] for x in (b.get("stages") or [])] if len(b.get("stages") or []) > 1 else []
            rows += [["Confident", f'{b["pct"]}%'], ["Not confident", f'{100 - b["pct"]}%']]
            _table(doc, ["Confidence", ""], rows, [64, 36], num=[1])
            _para(doc, b.get("note", ""), 8, color="5B6862", italic=True)
            if b.get("breakdown"):
                h = b["breakdown"]
                _table(doc, h["head"], h["rows"], small=True)
        elif t == "kpis":
            _table(doc, [k for _, k in b["items"]], [[v for v, _ in b["items"]]], num=list(range(len(b["items"]))))
        elif t == "table":
            _table(doc, b["head"], b["rows"], b.get("widths"), b.get("small", False), b.get("num"))
        elif t == "mono":
            tb = doc.add_table(rows=1, cols=1)
            c = tb.rows[0].cells[0]
            _shade(c, "F5F8F6")
            for i, line in enumerate(b["lines"]):
                p = c.paragraphs[0] if i == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                _run(p, line, 8.5, font="Consolas")
            doc.add_paragraph()
        elif t == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def safe_name(s: str) -> str:
    import re
    return re.sub(r"[^\w &.,()-]", "", s).strip().rstrip(".")
