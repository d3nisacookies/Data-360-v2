"""Fix plan: every change and open issue mapped back to the real uploaded file it lives in,
so data owners can correct the source. Built from the Clean audit trail, escalations,
invalid internal values, stale records and values the Seeker could not find."""
from __future__ import annotations

import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import ui
from .clean import business_days
from .reference import CRITICAL, LABEL
from .registry import meta
from .values import validate_element

SEV = ["Critical", "High", "Medium", "Low"]
TYPES = {"correct": "Correct a value", "dup": "Remove duplicate", "reconcile": "Align to golden value", "format": "Fix format or placeholder",
         "refresh": "Refresh stale record", "fill": "Add missing value", "investigate": "Investigate and confirm"}
TYPE_OF = {"Remove duplicates": "dup", "Correct errors": "correct", "Reconcile conflicting records": "reconcile", "Standardise formats": "format",
           "Harmonise currency and FX": "format", "Treat nulls explicitly": "fill", "Outlier and reasonableness testing": "investigate",
           "Meet data protection requirements": "format"}
LANES = [("Now", "This week", ("Critical", "High")), ("Next", "2–4 weeks", ("Medium",)), ("Later", "Within 6 weeks", ("Low",))]
CRIT_LABELS = {LABEL[f] for f in CRITICAL}


def build(raw, clean: dict, seek: dict | None, uploads: dict, owners: dict, resolved: set, ref: date) -> dict | None:
    """raw: the dataset Clean ran on (as sourced). uploads: {id: {"name", "folder"}}."""
    if not clean:
        return None
    F: dict[int, dict] = {}

    def ensure(upload_id, kind, source):
        if upload_id not in F:
            u = uploads.get(upload_id) or {"name": f"Upload {upload_id}", "folder": ""}
            F[upload_id] = {"k": str(upload_id), "upload_id": upload_id, "kind": kind, "name": u["name"], "folder": u.get("folder") or kind,
                            "owner": meta(source, kind, owners)["owner"] or "Unassigned", "sources": set(), "fixes": []}
        F[upload_id]["sources"].add(source)
        return F[upload_id]

    rec_by = {}
    for r in raw.records:
        rec_by.setdefault(r.rid, r)
    el_by = {e.key.split(":", 1)[-1]: e for e in raw.elements}
    first_rec_of = {}
    for r in raw.records:
        first_rec_of.setdefault(r.source, r)
    first_el_of = {}
    for e in raw.elements:
        first_el_of.setdefault(e.source, e)

    def owner(kind, src):
        return meta(src, kind, owners)["owner"] or "Unassigned"

    def push(file, **f):
        file["fixes"].append({"id": "", "due": "", **f})

    # 1 · Clean audit trail: applied in the cleaned pack, still to apply at source
    logged = set()
    for l in clean["log"]:
        r = rec_by.get(l["record"])
        e = None if r else el_by.get(l["record"])
        if r:
            file, kind = ensure(r.upload_id, "External", r.source), "External"
        elif e:
            file, kind = ensure(e.upload_id, "Internal", e.source), "Internal"
            logged.add(e.key)
        else:
            continue
        crit = l["field"] in CRIT_LABELS
        sev = "Critical" if re.search("sanction", l["field"], re.I) else "Low" if l["action"] == "Remove duplicates" else "High" if crit else "Medium"
        push(file, rec=l["record"], feed=l["source"], field=l["field"], **{"from": str(l["from"])}, to=str(l["to"]), why=l["reason"], type=TYPE_OF.get(l["action"], "correct"),
             sev=sev, owner=owner(kind, l["source"]), status="In data pack v1.1 · apply at source")
    # 2 · Escalations: investigate in the files named
    for ix, x in enumerate(clean["escalations"]):
        iid = f"ISS-{ix + 1:02d}"
        targets = []
        for rid in re.findall(r"[A-Z]{2,}-[\w-]+", x["item"]):
            if rid in rec_by:
                targets.append(ensure(rec_by[rid].upload_id, "External", rec_by[rid].source))
        if not targets:
            names = set(re.findall(r"\(([^()]+)\)\s*$", x["why"]) + re.findall(r"\(([^()]+)\)", x["item"]))
            m = re.search(r"Data owner · (.+)$", x["owner"])
            if m:
                names.add(m.group(1))
            m = re.search(r"from (.+)$", x["item"])
            if m:
                names.add(m.group(1))
            for n in {t.strip().split(" ")[0] if t.strip().split(" ")[0] in first_rec_of else t.strip() for n0 in names for t in n0.split(",")}:
                if n in first_rec_of:
                    targets.append(ensure(first_rec_of[n].upload_id, "External", n))
                elif n in first_el_of:
                    targets.append(ensure(first_el_of[n].upload_id, "Internal", n))
            for n in re.findall(r"([A-Z][\w &'/-]+?) (?:\S+)(?:;|$)", x["why"]):
                if n in first_el_of:
                    targets.append(ensure(first_el_of[n].upload_id, "Internal", n))
            if re.search("sanction", x["item"], re.I):
                for n in ("Legal systems", "CRM tools"):
                    if n in first_el_of:
                        targets.append(ensure(first_el_of[n].upload_id, "Internal", n))
        done = iid in resolved
        for file in {id(t): t for t in targets}.values():
            push(file, rec=next(iter(re.findall(r"[A-Z]{2,}-[\w-]+", x["item"])), "—"), feed="—", field=re.sub(r" for .*$", "", x["item"]), **{"from": "—"},
                 to="Confirm with owner", why=x["why"], type="investigate", due=x["due"],
                 sev="Critical" if re.search("sanction", x["item"], re.I) else "High" if re.search("registration|facility amount|outstanding|maturity", x["item"], re.I) else "Medium",
                 owner=x["owner"], status="Resolved (Phase 3)" if done else "Open")
    # 3 · invalid internal values
    for e in raw.elements:
        v = validate_element(e.dtype, e.value)
        if not v or e.key in logged:
            continue
        to = ("Real value from the system of record" if re.search("Placeholder|List of options", v) else "Add ISO currency code" if "currency" in v
              else "Add % or bps" if "unit" in v else "ISO date (YYYY-MM-DD)" if re.search("date", v, re.I) else "Valid value")
        push(ensure(e.upload_id, "Internal", e.source), rec=e.key.split(":", 1)[-1], feed=e.source, field=e.element, **{"from": str(e.value or "Blank")}, to=to, why=v,
             type="format", sev="Medium" if e.dtype in ("Amount", "Percentage", "Date") else "Low", owner=owner("Internal", e.source), status="Open")
    # 4 · stale external records
    removed = set(clean["changes"].get("removed", []))
    for r in raw.records:
        if r.key in removed or not raw.is_stale(r):
            continue
        age = f"{(ref - r.as_of).days} days old" if r.as_of else "Flagged stale by the provider"
        push(ensure(r.upload_id, "External", r.source), rec=r.rid, feed=r.source, field="Whole record", **{"from": f"As of {r.as_of or 'unknown'}"}, to=f"Latest record from {r.source}",
             why=age, type="refresh", sev="High" if re.search("sanction", r.source, re.I) else "Medium", owner=owner("External", r.source), status="Open")
    # 5 · values the memo needs that no file holds
    for a in (seek or {}).get("results", []):
        if a["status"] not in ("miss", "pending"):
            continue
        home = next((s for s in ("CRM tools", "Legal systems", "Loan systems") if s in first_el_of), next(iter(first_el_of), None))
        if not home:
            continue
        file = ensure(first_el_of[home].upload_id, "Internal", home)
        push(file, rec="New row", feed=home, field=a["attribute"], **{"from": "Not held"}, to=f"Add {a['attribute']} for {raw.borrower}",
             why="The Seeker found it in no source; the credit memo needs it", type="fill", sev="High" if re.search("LEI|Tax ID", a["attribute"]) else "Low",
             owner=owner("Internal", home), status="Open")
    files = [f for f in F.values() if f["fixes"]]
    days = {"Critical": 2, "High": 5, "Medium": 15, "Low": 30}
    for f in files:
        f["sources"] = sorted(f["sources"])
        f["fixes"].sort(key=lambda x: SEV.index(x["sev"]))
        for x in f["fixes"]:
            x["due"] = x["due"] or business_days(ref, days[x["sev"]])
        f["open"] = [x for x in f["fixes"] if not x["status"].startswith("Resolved")]
        f["worst"] = next((s for s in SEV if any(x["sev"] == s for x in f["open"])), "Done")
    files.sort(key=lambda f: ((SEV.index(f["worst"]) if f["worst"] in SEV else 9), -len(f["open"])))
    for fi, f in enumerate(files):
        for i, x in enumerate(f["fixes"]):
            x["id"] = f"F{fi + 1}-{i + 1:02d}"
    return {"c": raw.borrower, "files": files}


def counts(P):
    allx = [x for f in P["files"] for x in f["fixes"]]
    open_ = [x for x in allx if not x["status"].startswith("Resolved")]
    return allx, open_


def lane_files(P, sevs):
    return [{"f": f, "x": [x for x in f["open"] if x["sev"] in sevs]} for f in P["files"] if any(x["sev"] in sevs for x in f["open"])]


def view(P, file_key: str | None, typ: str = "all") -> dict:
    allx, open_ = counts(P)
    cnt = lambda s: sum(1 for x in open_ if x["sev"] == s)  # noqa: E731
    in_pack = sum(1 for x in allx if "data pack" in x["status"])
    sel = next((f for f in P["files"] if f["k"] == file_key), P["files"][0] if P["files"] else None)

    def tile(f):
        n = len(f["open"]) or 1
        return {"k": f["k"], "name": re.sub(r"\.(xlsx|csv|xls)$", "", f["name"], flags=re.I), "worst": f["worst"].lower(), "open": len(f["open"]),
                "done": len(f["fixes"]) - len(f["open"]), "mix": [{"s": s.lower(), "pct": sum(1 for x in f["open"] if x["sev"] == s) / n * 100} for s in SEV],
                "on": bool(sel and sel["k"] == f["k"])}
    groups = []
    for kind, label in (("Internal", "Internal source files"), ("External", "External data files")):
        fs = [f for f in P["files"] if f["kind"] == kind]
        if fs:
            groups.append({"label": label, "tiles": [tile(f) for f in fs]})
    lanes = []
    for n, w, sevs in LANES:
        fs = lane_files(P, sevs)
        lanes.append({"n": n, "w": w, "count": sum(len(o["x"]) for o in fs),
                      "cards": [{"k": o["f"]["k"], "name": re.sub(r"\.(xlsx|csv|xls)$", "", o["f"]["name"], flags=re.I),
                                 "what": f"{len(o['x'])} fix{'es' if len(o['x']) > 1 else ''}: " + ", ".join(dict.fromkeys(TYPES[x['type']].lower() for x in o["x"])),
                                 "due": min(x["due"] for x in o["x"])} for o in fs]})
    detail = None
    if sel:
        rows = [x for x in sel["fixes"] if typ == "all" or x["type"] == typ]
        detail = {"k": sel["k"], "name": sel["name"], "folder": sel["folder"], "sources": sel["sources"],
                  "types": [{"k": k, "t": t, "n": len(sel["fixes"]) if k == "all" else sum(1 for x in sel["fixes"] if x["type"] == k)} for k, t in [("all", "All"), *TYPES.items()]],
                  "type": typ,
                  "table": ui.table(["#", "Row / record", "Field", "Current", "Change to", "Why", "Severity", "Owner", "Due", "Status"],
                                    [[ui.mono(x["id"]), ui.bold(x["rec"], x["feed"] if x["feed"] != "—" else ""), ui.bold(x["field"], TYPES[x["type"]]), x["from"], ui.bold(x["to"]), ui.hint(x["why"]),
                                      ui.chip({"Critical": "bad", "High": "warn", "Medium": "info"}.get(x["sev"], "idle"), x["sev"]), ui.hint(x["owner"]), ui.mono(x["due"]),
                                      ui.chip("ok" if x["status"].startswith("Resolved") else "info" if "data pack" in x["status"] else "warn", x["status"])] for x in rows],
                                    row_class=["done" if x["status"].startswith("Resolved") else "" for x in rows])}
    return {"c": P["c"], "headline": f"**{len(open_)}** fixes across **{len(P['files'])}** files for **{P['c']}**. {in_pack} are already corrected in the cleaned data pack and need applying at source; "
                                      "the rest need investigation, refresh or new data.",
            "kpis": ui.kpis([(len(P["files"]), "Files to fix"), (cnt("Critical"), "Critical", "bad"), (cnt("High"), "High", "warn"), (cnt("Medium") + cnt("Low"), "Medium and low"),
                             (len(allx) - len(open_), "Resolved")], five=True),
            "groups": groups, "lanes": lanes, "detail": detail}


def xlsx(P) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="17644A")

    def sheet(ws, head, rows, widths):
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill = bold, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)].width = w
        ws.freeze_panes = "A2"
    sheet(ws, ["File", "Folder", "Sources", "Open fixes", *SEV, "Resolved"],
          [[f["name"], f["folder"], ", ".join(f["sources"]), len(f["open"]), *[sum(1 for x in f["open"] if x["sev"] == s) for s in SEV], len(f["fixes"]) - len(f["open"])] for f in P["files"]],
          [44, 26, 40, 10, 9, 7, 9, 7, 9])
    used = set()
    for i, f in enumerate(P["files"]):
        n = re.sub(r"[\\/?*\[\]:]", "", re.sub(r"\.(xlsx|csv|xls)$", "", f["name"], flags=re.I))[:28] or f"File {i + 1}"
        while n in used:
            n = n[:25] + str(i)
        used.add(n)
        sheet(wb.create_sheet(n), ["#", "Row / record", "Feed", "Field", "Fix type", "Current", "Change to", "Why", "Severity", "Owner", "Due", "Status"],
              [[x["id"], x["rec"], x["feed"], x["field"], TYPES[x["type"]], x["from"], x["to"], x["why"], x["sev"], x["owner"], x["due"], x["status"]] for x in f["fixes"]],
              [8, 14, 18, 26, 22, 26, 30, 44, 9, 24, 11, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
