"""Credit memo step 5 screens (5.1–5.5) as server-described sections, and the evidence pack."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import memo as MM
from . import ui
from .blocks import to_html
from .policies import POLICY

COVC = {"Compliant": "ok", "Tight": "warn", "Conflict": "bad", "Breach": "bad", "Not tested": "idle"}
RESC = {"Pass": "ok", "Fail": "bad", "Exception": "warn"}


def _sec_title(no):
    return "Document control" if no == "0" else (no + ". " if not re.fullmatch("[A-D]", no) else "") + MM.SEC[no]["t"]


def _conf_chip(r):
    if r["kind"] == "missing":
        return ui.chip("bad", "Not sourced")
    if r["kind"] in ("manual", "edited"):
        return ui.chip("info", "Reviewer" if r["kind"] == "edited" else "RM")
    return ui.chip({"High": "ok", "Medium": "warn", "Low": "bad"}.get(r["conf"], "idle"), r["conf"] or "—")


def setup(st) -> dict:
    sel = MM.selected(st)
    left = [s for s in MM.SECTIONS if not s[2] and s[0] not in st["extra"]]
    return {"core": [{"no": s[0], "t": s[1], "d": s[3], "on": s[0] in st["on"]} for s in MM.SECTIONS if s[2]],
            "extra": [{"no": no, "t": MM.SEC[no]["t"], "d": MM.SEC[no]["d"]} for no in st["extra"]],
            "pick": [{"no": s[0], "t": _sec_title(s[0])} for s in left], "count": f"{len(sel)} of {len(MM.SECTIONS)} sections",
            "stages": [{"k": i + 1, **s, "done": st["done"] >= i + 1} for i, s in enumerate(MM.STAGES)], "done": st["done"]}


def sections(M, ST, st, ref_date, retained_ok: bool, policies_on: int) -> list[dict]:
    out = []
    done = st["done"]
    sel = MM.selected(st)
    word = MM.WORD.get(M["capR"]["r"], "—")
    if done >= 1:
        rows = [[r["title"], r["n"], r["auto"], r["man"], r["miss"], ui.bar(round(r["auto"] / r["n"] * 100) if r["n"] else 0, "", f"{round(r['auto'] / r['n'] * 100) if r['n'] else 0}%"),
                 ui.hint(", ".join(r["srcs"]))] for r in ST["bySec"]]
        rows += [[_sec_title(no), {"t": "span", "text": "Built from Analyse, Compliance and data-quality results", "cls": "hint", "colspan": 6}]
                 for no in sel if not any(r["no"] == no for r in ST["bySec"])]
        out.append(ui.section("5.1 Prepare", M["c"], [
            ui.headline(f"**{ST['total']}** template fields in scope for **{M['c']}**: **{ST['auto']}** filled automatically, **{ST['manual'] + ST['edited']}** for the RM or reviewer, "
                        f"**{len(ST['missing'])}** not sourced."),
            ui.kpis([(ST["total"], "Fields in scope"), (ST["auto"], "Filled automatically"), (ST["manual"] + ST["edited"], "RM or reviewer"), (len(ST["missing"]), "Not sourced"),
                     (ST["srcs"], "Sources used")], five=True),
            ui.table(["Section", "Fields", "Auto", "RM", "Not sourced", "Coverage", "Main sources"], rows, num=[1, 2, 3, 4]),
            ui.details(f"Not sourced · {len(ST['missing'])} fields go to the Exception Register",
                       [ui.table(["Section", "Field", "Why"], [["Control" if f["sec"] == "0" else f["sec"], f["label"], ui.hint(f["note"])] for f in ST["missing"]])])], id="cm1"))
    if done >= 2:
        reds = sum(1 for r in M["RAG"] if r["rag"] == "r")
        L = ["Current" if i == len(M["mLabels"]) - 1 else m for i, m in enumerate(M["mLabels"])]
        carry = M["carry"]
        out.append(ui.section("5.2 Analyse", "Ratios, trend, covenants, risk ratings", [
            ui.headline(f"Net leverage **{M['levTxt']}**, LTV **{M['ltvTxt']}**, " + (f"pricing **{'+' if carry >= 0 else '−'}{abs(carry)} bps** vs cost of funds, " if carry is not None else "")
                        + f"**{reds}** of {len(M['RAG'])} risk areas red."),
            ui.h3("Historical performance"), ui.table(["Metric", *L, "Trend"], [[r[0], *r[1], r[2]] for r in M["finRows"]], num=list(range(1, len(L) + 1)), flat=False),
            ui.hint_block("Audited annual accounts are not held, so the trend uses the latest validated observation per month."),
            ui.h3("Risk assessment"), ui.table(["Risk area", "Assessment", "RAG"], [[ui.bold(r["area"]), r["txt"], ui.chip(*ui.RAGC[r["rag"]])] for r in M["RAG"]], flat=False),
            ui.h3("Covenants"), ui.table(["Covenant", "Threshold", "Actual", "Headroom", "Status", "Note"],
                                         [[ui.bold(r["name"]), ui.mono(r["thrTxt"]), ui.mono(r["actTxt"]), ui.mono(r["headTxt"]), ui.chip(COVC.get(r["status"], "idle"), r["status"]), ui.hint(r["note"])]
                                          for r in M["COV"]], flat=False),
            {"type": "twocol", "cols": [{"title": "Strengths", "items": M["strengths"] or ["None identified"]}, {"title": "Concerns", "items": M["concerns"] or ["None identified"]}]}], id="cm2"))
    if done >= 3:
        n = lambda r: sum(1 for x in M["CH"] if x["result"] == r)  # noqa: E731
        tone = "ok" if M["decision"] == "Approve" else "bad" if M["decision"] == "Decline" else ""
        blocks = [{"type": "decision", "tone": tone, "label": "Draft decision", "value": M["decision"], "why": M["decWhy"]},
                  ui.kpis([(n("Pass"), "Pass", "ok"), (n("Fail"), "Fail", "bad"), (n("Exception"), "Exception", "warn"), (M["hardFails"], "Hard failures (refer)")])]
        if M["CHoff"]:
            blocks.append({"type": "polnote", "text": f"{len(M['CHoff'])} check{'s are' if len(M['CHoff']) > 1 else ' is'} switched off in Rules & policies and not applied: "
                                                      + "; ".join(f"{x['check']} ({x['rule']})" for x in M["CHoff"]) + "."})
        blocks += [ui.table(["Area", "Rule", "Check", "Result", "Evidence", "Source"],
                            [[ui.hint(x["area"]), ui.cell(x["rule"] or "—", "mono", f"{POLICY[x['rule']]['t']} · {POLICY[x['rule']]['src']}" if x["rule"] in POLICY else ""),
                              ui.bold(x["check"] + (" · hard" if x["hard"] and x["result"] == "Fail" else "")), ui.chip(RESC[x["result"]], x["result"]), x["evidence"], ui.hint(x["src"])]
                             for x in M["CH"]]),
                   {"type": "twocol", "ordered": True, "cols": [{"title": f"Conditions precedent ({len(M['CP'])})", "items": M["CP"]},
                                                                {"title": f"Conditions subsequent ({len(M['CS'])})", "items": M["CS"]}]}]
        out.append(ui.section("5.3 Rate against compliance", f"{len(M['CH'])} checks · {policies_on} rules active", blocks, id="cm3"))
    if done >= 4:
        B = MM.blocks(M, ST, st, ref_date)
        TR = MM.trace_rows(M, ST, st)
        out.append(ui.section("5.4 Present", "Credit Memo Acceleration & Assurance template", [
            ui.kpis([(M["ref"], f"Reference · {MM.version(st)}"), (f"{M['last'].get('pct')}%", f"Data confidence ({word})"), (f"{ST['ai']}%", "AI confidence"),
                     (f"{ST['pctAuto']}%", "AI-generated content"), (M["decision"], "Draft decision")], five=True),
            ui.details(f"Preview memo · {len(sel)} sections", [{"type": "html", "html": to_html(B)}]),
            ui.details(f"Traceability · {len(TR)} values with source and confidence", [
                {"type": "filter", "placeholder": "Filter by field, value or source"},
                ui.table(["Sec", "Field", "Value", "Source", "Confidence"],
                         [[ui.hint(r["sec"]), ui.bold(r["label"], r["note"]), r["v"], ui.hint(r["src"] + (f" · {r['when']}" if r["when"] else "")), _conf_chip(r)] for r in TR])])],
            foot={"hint": "Generates a new Word file. Nothing in your source files changes.",
                  "downloads": [{"label": "Download HTML", "href": "memo.html"}, {"label": "Generate credit memo (.docx)", "href": "memo.docx", "primary": True}]}, id="cm4"))
    if done >= 5:
        R = st.get("retained")
        ap = st["approvals"]
        nxt = next((i for i, a in enumerate(ap) if not a["decision"]), -1)
        fs = MM.fields_in(M, ["0", *sel])
        blocks = [ui.h3("Approval chain"), ui.hint_block(f"Each approver decides in turn. The draft decision from 5.3 is {M['decision']}."),
                  {"type": "approvals", "locked": bool(R), "next": nxt, "decisions": MM.DECISIONS,
                   "rows": [{"i": i, **a, "can": not R and (i == 0 or bool(ap[i - 1]["decision"]))} for i, a in enumerate(ap)]},
                  ui.h3("Reviewer edits"),
                  ui.hint_block("Change a value in the memo. Each edit is logged and counts as manual content. Compliance checks still use the source data."),
                  {"type": "edit", "disabled": bool(R), "fields": [{"key": f["key"], "label": ("Control" if f["sec"] == "0" else f["sec"]) + " · " + f["label"], "v": f["v"]} for f in fs]}]
        E = st.get("edits") or []
        blocks.append(ui.table(["#", "Field", "Was", "Now", "Reason", "By", "When"], [[i + 1, e["label"], e["from"], e["to"], e["reason"], e["by"], e["at"]] for i, e in enumerate(E)], flat=False)
                      if E else ui.hint_block("No edits yet."))
        blocks.append(ui.h3("Retain"))
        if R:
            blocks += [{"type": "decision", "tone": "ok", "label": "Retained", "value": f"{R['version']} · {MM.status(st)}", "why": f"{R['at']} · {R['location']}"},
                       ui.table(["Item", "Value"], [["SHA-256 fingerprint", ui.mono(R["hash"])], ["Retention", R["retention"]], ["Legal hold", R["legalHold"]],
                                                    ["Evidence", "Memo, Appendices A–D, compliance checks, approvals, edits"]], flat=False)]
        else:
            waiting = sum(1 for a in ap if not a["decision"])
            blocks.append(ui.hint_block("All approvers have decided. Finalising locks the memo, stamps the version and records a fingerprint." if ap[-1]["decision"]
                                        else f"Waiting for {waiting} approver{'' if waiting == 1 else 's'}. Finalising unlocks after the Credit Committee decides."))
        foot = {"hint": "Final files are new copies. Your source files are never changed." if R else "The evidence pack is available once the memo is retained."}
        if R:
            foot["downloads"] = [{"label": "Export evidence pack (.xlsx)", "href": "evidence.xlsx"}, {"label": "Download final memo (.docx)", "href": "memo.docx", "primary": True}]
        else:
            foot["buttons"] = [ui.btn("Finalise and retain", "retain", None, disabled=not (ap[-1]["decision"] and retained_ok), primary=True)]
        out.append(ui.section("5.5 Review and retain", f"{MM.status(st)} · {MM.version(st)}", blocks, foot=foot, id="cm5"))
    return out


def evidence_xlsx(M, ST, st, ref_date) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    hf, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="17644A")
    R = st.get("retained")

    def add(name, rows, widths):
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        for c in ws[1]:
            c.font, c.fill = hf, fill
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
    B = MM.blocks(M, ST, st, ref_date)
    add("Memo control", [["Item", "Details"], *B[0]["meta"], ["SHA-256 fingerprint", R["hash"] if R else "Not retained"], ["Retention", R["retention"] if R else "—"]], [28, 90])
    add("A Source traceability", [["Section", "Data element", "Value", "Source", "Timestamp", "Owner", "Confidence", "Note"],
                                  *[[r["sec"], r["label"], r["v"], r["src"], r["when"], r["owner"], "Not sourced" if r["kind"] == "missing" else r["conf"] or "", r["note"]] for r in MM.trace_rows(M, ST, st)]],
        [8, 30, 50, 34, 12, 26, 12, 50])
    rows = []
    for d, s, g in MM.DOMAINS:
        r = MM.dom_score(M, s, ref_date) or {}
        rows.append([d, g, r.get("Completeness"), r.get("Accuracy"), r.get("Consistency"), r.get("Validity"), r.get("Timeliness"), r.get("overall")])
    add("B DQ dashboard", [["Domain", "Golden source", "Completeness", "Accuracy", "Consistency", "Validity", "Timeliness", "Overall"], *rows], [14, 26, 13, 11, 12, 10, 11, 10])
    add("C Reconciliation log", [["Attribute", "Sources compared", "Golden value", "Status"], *MM.recon_rows(M)], [24, 50, 30, 14])
    add("D Assurance", [["Area", "Status"], *MM.final_conf(M, ref_date)], [30, 50])
    add("Compliance checks", [["Area", "Rule", "Check", "Result", "Hard", "Evidence", "Source", "Condition"],
                              *[[x["area"], x["rule"] or "", x["check"], x["result"], "Y" if x["hard"] else "N", x["evidence"], x["src"], x["cp"] or x["cs"] or ""] for x in M["CH"]]],
        [14, 8, 32, 10, 6, 70, 26, 50])
    add("Exceptions", [["Exception", "Impact", "Resolution"], *[[e["ex"], e["impact"], e["res"]] for e in MM.exceptions(M, ST)]], [80, 10, 60])
    add("Approvals", [["Approver", "Name", "Decision", "Comment", "Date"], *[[a["role"], a["name"], a["decision"], a["comment"], a["date"]] for a in st["approvals"]]], [22, 22, 24, 40, 12])
    add("Reviewer edits", [["#", "Field", "Was", "Now", "Reason", "By", "When"], *[[i + 1, e["label"], e["from"], e["to"], e["reason"], e["by"], e["at"]] for i, e in enumerate(st.get("edits") or [])]],
        [4, 30, 30, 30, 40, 20, 20])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
