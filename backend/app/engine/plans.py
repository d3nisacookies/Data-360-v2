"""Project plans for each Phase 3 stage: workstreams and tasks built from the actual issues,
datasets, fixes and roadmap items, scheduled in business days from the stage start (a Monday)."""
from __future__ import annotations

import io
import math
import re
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import phase3 as P3
from . import ui
from .fixplan import lane_files
from .phase2 import horizon_of, usd
from .registry import meta

STAGE_START = {"stab": 0, "std": 3, "opt": 6, "trans": 12}
STAGE_END = {"stab": 3, "std": 6, "opt": 12, "trans": 18}


def bd_add(d0: date, n: int) -> date:
    d, k = d0, 0
    while k < n:
        d += timedelta(days=1)
        if d.weekday() < 5:
            k += 1
    return d


def stage_start(ref: date, k: str) -> date:
    m = ref.month - 1 + STAGE_START[k]
    y, m = ref.year + m // 12, m % 12 + 1
    d = date(y, m, min(ref.day, 28))
    while d.weekday() != 0:
        d += timedelta(days=1)
    return d


def build(k: str, ctx: dict, st: dict, fix: dict | None) -> dict:
    S = P3.STEP[k]
    p2 = ctx["p2"]
    R = p2["R"]["road"]
    recs = {r["id"]: r for r in p2["R"]["rec"]["R"]}
    rate = p2["assume"]["rate"] * 8
    owners = ctx.get("owners") or {}
    T, WS = [], []

    def ws(code, name, owner):
        WS.append({"code": code, "name": name, "owner": owner})
        return code

    def task(w, t, deliv, owner, dur, fte, dep=None, after=None, ms=False, link="", status=None, lanes=None):
        same = [x for x in T if x["ws"] == w]
        tid = f"{w}.{len(same) + 1}"
        if lanes and len(same) >= lanes:
            pool = [x for x in same if not x["ms"]]
            if len(pool) >= lanes:
                dep = pool[len(pool) - lanes]["id"]
        T.append({"id": tid, "ws": w, "t": t, "deliv": deliv, "owner": owner, "dur": dur, "fte": fte, "dep": dep, "after": after, "ms": ms, "link": link, "st": status})
        return tid

    def recs_in(h):
        return [{**recs[rid], "init": i["name"]} for i in R["I"] for rid in i["recs"] if horizon_of(recs[rid]) == h]

    def fix_lane(sevs):
        return lane_files(fix, sevs) if fix else []

    def fname(f):
        return re.sub(r"\.(xlsx|csv|xls)$", "", f["name"], flags=re.I)

    w0 = ws("PM", "Programme management", "Transformation office")
    kick = task(w0, f"{S['t']} kick-off and scope confirmation", "Stage charter signed", "Chief Data Officer", 2, 1, ms=True)
    if k == "stab":
        task(w0, "Stand up the data council", "Charter, members, cadence", "Chief Data Officer", 5, 0.5, dep=kick)
        w1 = ws("INV", "Data inventory", "Data Governance Office")
        S_ = P3.sources(ctx)
        prev = task(w1, f"Validate inventory: {len(S_['int'])} internal systems", "Confirmed element counts and refresh cycles", "Data Governance Office", 5, 1, dep=kick)
        prev = task(w1, f"Validate inventory: {len(S_['ext'])} external feeds", "Confirmed feeds, methods and licences", "Market Data Services", 4, 1, dep=prev)
        task(w1, "Publish data catalogue v1", "Catalogue live", "Data Governance Office", 3, 0.5, dep=prev, ms=True)
        w2 = ws("OWN", "Ownership", "Data Governance Office")
        for x in S_["int"] + S_["ext"]:
            if meta(x["key"], x["kind"], {})["owner"]:
                continue
            got = owners.get(x["key"])
            task(w2, f"Assign owner and steward: {x['name']}", "Named owner and steward", "Data Governance Office", 5, 0.25, dep=kick,
                 link=f"{x['kind']} · assigned to {got}" if got else x["kind"], lanes=2, status="Done" if got else None)
        task(w2, "Ownership register sign-off", "RACI approved by data council", "Chief Data Officer", 2, 0.25, after="OWN", ms=True)
        w3 = ws("ISS", "Critical issue resolution", "Financial Crime Compliance / Credit Risk")
        for i in P3.issues(ctx):
            if i["sev"] in ("Critical", "High"):
                task(w3, f"{i['id']}: {i['t']}", "Issue closed with evidence", i["owner"], 5 if i["sev"] == "Critical" else 8, 0.5, dep=kick, link=i["sev"], lanes=4,
                     status="Done" if i["id"] in st["resolved"] else "Not started")
        w4 = ws("FIX", "Source file fixes (Now lane)", "Data owners")
        for o in fix_lane(("Critical", "High")):
            task(w4, f"Apply {len(o['x'])} fix{'es' if len(o['x']) > 1 else ''} to {fname(o['f'])}", "Corrected source file, re-run Clean", o["f"]["owner"],
                 min(8, max(3, len(o["x"]))), 0.5, dep=kick, lanes=3, link=", ".join(x["id"] for x in o["x"][:4]))
        w5 = ws("QW", "Quick wins from the roadmap", "Initiative owners")
        for r in recs_in(0):
            task(w5, f"{r['id']}: {r['t']}", r["out"], r["owner"], 15 if r["e"] == "S" else 30, 0.5, dep=kick, link=r["init"], lanes=3)
        w6 = ws("BASE", "Executive confidence baseline", "Data council")
        b1 = task(w6, "Freeze confidence baseline by domain", f"Baseline {st['baseline']['base'] if st.get('baseline') else '—'}% recorded", "Data Governance Office", 2, 0.5, after="ISS")
        task(w6, "Executive readout: Stabilise outcomes", "Readout pack to CRO and Head of Credit", "Chief Data Officer", 2, 0.5, dep=b1, ms=True)
    if k == "std":
        w1 = ws("DS", "Trusted datasets", "Data Governance Office")
        for d in P3.datasets(ctx, st):
            a = task(w1, f"{d['id']} {d['name']}: build from cleaned pack", f"{len(d['fields'])} golden fields loaded", "Data engineering", 5, 1, dep=kick)
            b = task(w1, f"{d['id']} {d['name']}: quality tests", "Rules pass at 100% on critical fields", "Data Governance Office", 3, 0.5, dep=a)
            task(w1, f"{d['id']} {d['name']}: certify", "Certified dataset", "Data council", 1, 0.25, dep=b, ms=True, status="Done" if d["id"] in st["certified"] else "Not started")
        w2 = ws("KPI", "Consistent reporting and KPIs", "Credit Risk Management")
        K = (ctx["norm"] or {}).get("out", {}).get("KPIs", {}).get("rows", [])
        kp = task(w2, f"Publish KPI catalogue ({len(K)} KPIs)", "KPI definitions, owners and thresholds", "Data Governance Office", 5, 0.5, dep=kick)
        for r in ("Credit memo pack", "Data quality scorecard", "Escalation report", "Portfolio KPI pack"):
            kp = task(w2, f"Build standard report: {r}", "Report on certified definitions", "BI team", 5, 1, dep=kp)
        w3 = ws("MAN", "Reduce manual effort", "Credit analyst")
        m = task(w3, "Map memo fields to certified datasets", "Field-to-dataset map", "Credit analyst", 5, 0.5, after="DS")
        m = task(w3, "Auto-populate memo fields from datasets", "Memo template wired to datasets", "Technology", 10, 1, dep=m)
        task(w3, "Measure hours per memo (before and after)", f"Target {R['hNow']} h → lower", "Credit analyst", 3, 0.25, dep=m, ms=True)
        w4 = ws("DEF", "Definitions and data contracts", "Data council")
        d1 = task(w4, "Certify common definitions v1.1", "Approved definitions register", "Data council", 5, 0.5, dep=kick, ms=True)
        for x in P3.sources(ctx)["ext"]:
            task(w4, f"Data contract: {x['name']}", "Signed SLA and quality terms", "Procurement", 10, 0.25, dep=d1, lanes=5)
        w5 = ws("REC", "Roadmap items (3–6 months)", "Initiative owners")
        for r in recs_in(1):
            task(w5, f"{r['id']}: {r['t']}", r["out"], r["owner"], {"S": 15, "M": 30}.get(r["e"], 45), 1, dep=kick, link=r["init"], lanes=4)
        w6 = ws("FIX", "Source file fixes (Next lane)", "Data owners")
        for o in fix_lane(("Medium",)):
            task(w6, f"Apply {len(o['x'])} fix{'es' if len(o['x']) > 1 else ''} to {fname(o['f'])}", "Corrected source file", o["f"]["owner"], min(10, max(3, len(o["x"]))), 0.5, dep=kick, lanes=3)
    if k == "opt":
        w1 = ws("RPT", "Automated reporting", "BI team")
        r1 = kick
        for r in ("Escalation report (daily)", "Data quality scorecard (weekly)", "Portfolio KPI pack (monthly)", "Credit memo pack (on demand)"):
            r1 = task(w1, f"Automate: {r}", "Scheduled, monitored job", "BI team", 8, 1, dep=r1)
        task(w1, "Retire manual report versions", "Manual reports switched off", "Credit Risk Management", 5, 0.5, dep=r1, ms=True)
        w2 = ws("ANA", "Advanced analytics", "Credit Risk Management")
        a = task(w2, "Source reliability analytics", "Error-rate trend by feed", "Data analytics", 10, 1, dep=kick)
        a = task(w2, "Financial trend and covenant headroom analytics", "Headroom alerts by borrower", "Data analytics", 10, 1, dep=a)
        task(w2, f"Portfolio benchmark across {ctx.get('n_borrowers', 1)} borrowers", "Benchmark dashboard", "Data analytics", 10, 1, dep=a, ms=True)
        w3 = ws("VIS", "Visibility (control tower)", "Data Governance Office")
        v = task(w3, "Design control tower", "Wireframes approved", "BI team", 5, 0.5, dep=kick)
        v = task(w3, "Build control tower (confidence, issues, datasets, KPIs)", "Live control tower", "BI team", 15, 1, dep=v)
        task(w3, "Go-live and training", "Users trained", "Data Governance Office", 5, 0.5, dep=v, ms=True)
        w4 = ws("DEC", "Cross-functional decision support", "Chief Data Officer")
        for f in ("Credit", "Risk", "Compliance", "Treasury", "Operations"):
            task(w4, f"Decision view: {f}", f"{f} view on certified data", "BI team", 12, 0.5, after="VIS", lanes=2)
        w5 = ws("REC", "Roadmap items (6–12 months)", "Initiative owners")
        for r in recs_in(2):
            task(w5, f"{r['id']}: {r['t']}", r["out"], r["owner"], {"S": 20, "M": 40}.get(r["e"], 60), 1, dep=kick, link=r["init"], lanes=4)
        w6 = ws("FIX", "Source file fixes (Later lane)", "Data owners")
        for o in fix_lane(("Low",)):
            task(w6, f"Apply {len(o['x'])} fix{'es' if len(o['x']) > 1 else ''} to {fname(o['f'])}", "Corrected source file", o["f"]["owner"], min(10, max(3, math.ceil(len(o["x"]) / 2))), 0.25, dep=kick, lanes=3)
    if k == "trans":
        w1 = ws("AI", "AI-assisted memo drafting", "Data & AI office")
        a = task(w1, f"Set drafting threshold (confidence ≥ {st['aiThr']}) and guardrails", "Approved AI use policy for memos", "Data & AI office", 5, 0.5, dep=kick)
        a = task(w1, "Pilot on one portfolio", "Pilot results: time saved, accuracy", "Credit analyst", 20, 1, dep=a)
        a = task(w1, "Model risk review of drafting assistant", "Independent review sign-off", "Model risk management", 10, 0.5, dep=a)
        task(w1, "Scale to all memo types", "Drafting live for all memos", "Technology", 20, 1, dep=a, ms=True)
        w2 = ws("EWS", "Predictive early warning", "Credit Risk Management")
        e = task(w2, "Assemble default and downgrade history", "Modelling dataset", "Data analytics", 15, 1, dep=kick)
        e = task(w2, "Fit and calibrate early-warning model", "Calibrated model (replaces hand-set weights)", "Data analytics", 20, 1, dep=e)
        e = task(w2, "Independent model validation", "Validation report", "Model risk management", 15, 0.5, dep=e)
        task(w2, "Deploy to watchlist process", "Weekly early-warning run", "Technology", 10, 1, dep=e, ms=True)
        w3 = ws("OPT", "Continuous optimisation", "Data council")
        for x in ("Weekly data quality drift check", "Monthly confidence refresh against baseline", "Quarterly model back-test", "Six-monthly definition re-certification"):
            task(w3, f"Set up: {x}", "Running cadence with owner", "Data Governance Office", 5, 0.25, dep=kick)
        w4 = ws("ALR", "Proactive decisions", "Credit Risk Management")
        for t in dict.fromkeys(re.sub(r"\s*\(.*\)$", "", a2["t"]) for a2 in P3.alerts(ctx)):
            task(w4, f"Automate alert: {t}", "Alert with owner and action", "Technology", 10, 0.5, dep=kick, lanes=2)
        task(w4, "Alert playbooks signed off", "Playbook per alert", "Credit Risk Management", 5, 0.5, after="ALR", ms=True)
        w5 = ws("REC", "Roadmap items (12+ months)", "Initiative owners")
        for r in recs_in(3):
            task(w5, f"{r['id']}: {r['t']}", r["out"], r["owner"], 80 if r["e"] == "L" else 40, 1, dep=kick, link=r["init"], lanes=2)
    task(w0, f"{S['t']} close-out and sign-off", "Stage report and lessons learned", "Chief Data Officer", 3, 0.5, after="*", ms=True)
    # schedule
    start = stage_start(ctx["ref_date"], k)
    by = {x["id"]: x for x in T}
    for x in T:
        if x["after"] == "*":
            deps = [y for y in T if y is not x]
        elif x["after"]:
            deps = [y for y in T if y["ws"] == x["after"] and y is not x]
        else:
            deps = [by[x["dep"]]] if x["dep"] else []
        s0 = max([d["endIdx"] for d in deps if "endIdx" in d] + [0])
        x["startIdx"], x["endIdx"] = s0, s0 + x["dur"]
        x["start"] = start if s0 == 0 else bd_add(start, s0)
        x["end"] = start if x["endIdx"] - 1 <= 0 else bd_add(start, max(s0, x["endIdx"] - 1))
        x["effort"] = round(x["dur"] * x["fte"] * 10) / 10
        x["cost"] = round(x["effort"] * rate)
        x["st"] = x["st"] or "Not started"
        x["deps"] = "All tasks" if x["after"] == "*" else f"All {x['after']} tasks" if x["after"] else x["dep"] or "—"
        x["start"], x["end"] = x["start"].isoformat(), x["end"].isoformat()
    WS = [w for w in WS if any(x["ws"] == w["code"] for x in T)]
    end_idx = max(x["endIdx"] for x in T)
    risk3 = ("Model risk approval takes longer than planned" if k == "trans" else "Report consumers keep using manual versions" if k == "opt"
             else "Datasets blocked by open issues" if k == "std" else "Critical issues need third parties (borrower, vendors)")
    return {"k": k, "n": S["n"], "t": S["t"], "m": S["m"], "c": ctx["borrower"], "T": T, "WS": WS, "start": start.isoformat(), "end": bd_add(start, end_idx - 1).isoformat(),
            "weeks": math.ceil(end_idx / 5), "windowWeeks": round((STAGE_END[k] - STAGE_START[k]) * 4.33), "effort": round(sum(x["effort"] for x in T)),
            "cost": sum(x["cost"] for x in T), "ms": [x for x in T if x["ms"]], "rate": rate, "hourly": p2["assume"]["rate"],
            "risks": [["Owner capacity alongside business-as-usual", "High", "Ring-fence named owners for the stage; weekly stand-up"],
                      ["Source fixes reopen after re-runs", "Medium", "Re-run Collector and Clean after each batch of fixes"],
                      [risk3, "Medium", "Escalate through the data council at the fortnightly checkpoint"]]}


def section(P) -> dict:
    W = max(P["weeks"], 1)
    over = P["weeks"] > P["windowWeeks"]
    step = 4 if W > 20 else 2 if W > 10 else 1
    rows = []
    for w in P["WS"]:
        rows.append({"ws": w["name"]})
        for x in P["T"]:
            if x["ws"] == w["code"]:
                rows.append({"label": f"{x['id']} {x['t']}", "left": x["startIdx"] / (W * 5) * 100, "width": max(x["dur"], 0 if x["ms"] else 1) / (W * 5) * 100,
                             "ms": x["ms"], "done": x["st"] == "Done", "title": f"{x['t']}: {x['start']} → {x['end']}"})
    return ui.section(f"Project plan · {P['n']} {P['t']}", f"{P['c']} · {P['start']} to {P['end']}", [
        ui.kpis([(len(P["T"]), f"Tasks in {len(P['WS'])} workstreams"), (f"{P['weeks']} wk", f"Duration (window {P['windowWeeks']} wk)"), (len(P["ms"]), "Milestones"),
                 (P["effort"], "Person-days"), (usd(P["cost"]), "Internal effort cost")], five=True),
        ui.hint_block(f"The plan runs {P['weeks'] - P['windowWeeks']} weeks past the {P['m']} window. Add people to the longest workstream or move roadmap items to the next stage." if over
                      else f"Fits inside the {P['m']} window.", "warn" if over else "ok"),
        ui.h3("Timeline"),
        {"type": "pgantt", "weeks": [f"W{i + 1}" if i % step == 0 else "" for i in range(W)], "rows": rows},
        ui.hint_block(f"Bars show business days from {P['start']}. Diamonds are milestones. Green bars are already done."),
        ui.h3("Tasks"),
        ui.table(["ID", "Task", "Deliverable", "Owner", "Start", "End", "Days", "Effort", "Depends on", "Status"],
                 [[ui.mono(x["id"]), ui.bold(("◆ " if x["ms"] else "") + x["t"], x["link"]), ui.hint(x["deliv"]), ui.hint(x["owner"]), ui.mono(x["start"]), ui.mono(x["end"]),
                   x["dur"], x["effort"], ui.mono(x["deps"]), ui.chip("ok" if x["st"] == "Done" else "idle", x["st"])] for x in P["T"]], num=[6, 7], flat=False),
        ui.h3("Milestones"), ui.table(["Milestone", "Date", "Owner"], [[x["t"], x["end"], x["owner"]] for x in P["ms"]], num=[1], flat=False),
        ui.h3("Risks"), ui.table(["Risk", "Rating", "Mitigation"], P["risks"], flat=False)],
        foot={"hint": f"Effort cost uses USD {P['hourly']}/hour from the Phase 2 business case. Durations are planning estimates.",
              "download": {"label": "Download project plan (.xlsx)", "href": f"plan/{P['k']}.xlsx"}}, id="plan")


def xlsx(P) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    hf, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="17644A")

    def add(name, rows, widths):
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font, c.fill = hf, fill
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
    add("Overview", [["Item", "Value"], ["Client", P["c"]], ["Stage", f"{P['n']} {P['t']} ({P['m']})"], ["Start", P["start"]], ["End", P["end"]], ["Duration (weeks)", P["weeks"]],
                     ["Window (weeks)", P["windowWeeks"]], ["Tasks", len(P["T"])], ["Workstreams", len(P["WS"])], ["Milestones", len(P["ms"])], ["Effort (person-days)", P["effort"]],
                     ["Internal effort cost (USD)", P["cost"]], ["Day rate (USD)", P["rate"]]], [28, 40])
    wsn = {w["code"]: w["name"] for w in P["WS"]}
    add("Plan", [["ID", "Workstream", "Task", "Deliverable", "Owner", "Start", "End", "Business days", "FTE", "Effort (person-days)", "Cost (USD)", "Depends on", "Milestone", "Link", "Status"],
                 *[[x["id"], wsn[x["ws"]], x["t"], x["deliv"], x["owner"], x["start"], x["end"], x["dur"], x["fte"], x["effort"], x["cost"], x["deps"], "Yes" if x["ms"] else "", x["link"], x["st"]] for x in P["T"]]],
        [8, 26, 52, 36, 26, 11, 11, 8, 6, 10, 10, 14, 9, 24, 12])
    W = P["weeks"]

    def mark(x, i):
        s, e = i * 5, i * 5 + 5
        if x["ms"] and s <= x["endIdx"] - 1 < e:
            return "◆"
        return "■" if x["startIdx"] < e and x["endIdx"] > s else ""
    add("Gantt", [["ID", "Task", *[f"W{i + 1}" for i in range(W)]], *[[x["id"], x["t"], *[mark(x, i) for i in range(W)]] for x in P["T"]]], [8, 52, *[4] * W])
    add("Milestones", [["Milestone", "Date", "Owner"], *[[x["t"], x["end"], x["owner"]] for x in P["ms"]]], [56, 12, 26])
    add("Workstreams", [["Code", "Workstream", "Lead", "Tasks", "Effort (person-days)"],
                        *[[w["code"], w["name"], w["owner"], sum(1 for x in P["T"] if x["ws"] == w["code"]), round(sum(x["effort"] for x in P["T"] if x["ws"] == w["code"]) * 10) / 10] for w in P["WS"]]],
        [8, 34, 34, 8, 12])
    add("Risks", [["Risk", "Rating", "Mitigation"], *P["risks"]], [50, 10, 60])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
